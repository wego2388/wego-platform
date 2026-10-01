"""Bootstrap a new Safari Tours Sharm owner-data workbook.

This script intentionally generates only the original catalogue-backed base
sheets. The checked-in workbook is the working source of truth and contains
later owner answers plus account-audit and marketing sheets. To avoid silent
data loss, an existing workbook is never overwritten without the explicit
``--force-bootstrap`` flag.

Never put passwords, API keys or secret tokens in the workbook. See
OWNER_DATA_HUB.md for the live schema and agent rules.
"""
import argparse
import json
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
RESEARCH = HERE.parent / "content-research"
OUT = HERE / "safari-tours-owner-data-hub.xlsx"

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="0A2342")
EDIT = PatternFill("solid", fgColor="FFF8DC")   # owner fills: light yellow
LOCK = PatternFill("solid", fgColor="E7E6E6")   # system / do not edit: grey
DONE = PatternFill("solid", fgColor="E2F0D9")   # answered / filled: light green
FLAG = PatternFill("solid", fgColor="FDE9E7")   # needs a decision: light red
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS = ["ناقص", "اتملى", "محتاج مراجعة", "مش محتاجينه"]
PRIORITY = ["ضروري للإطلاق", "مهم بعد الإطلاق", "اختياري"]


def sheet(wb, title, columns, rows, editable, widths, status_col=None, priority_col=None, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for row in rows:
        ws.append(row)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(columns) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            name = columns[c - 1]
            if name in editable:
                cell.fill = DONE if cell.value not in (None, "") else EDIT
            else:
                cell.fill = LOCK
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 34
    for col_name, values in ((status_col, STATUS), (priority_col, PRIORITY)):
        if col_name:
            letter = openpyxl.utils.get_column_letter(columns.index(col_name) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{max(ws.max_row, 2) + 50}")
    return ws


def readme(wb):
    ws = wb.active
    ws.title = "اقرأني"
    ws.sheet_view.rightToLeft = True
    lines = [
        ("ملف بيانات Safari Tours Sharm — كل حاجة في مكان واحد", True),
        (f"اتعمل {date.today().isoformat()} — المصدر الواحد لكل البيانات اللي الموقع والـERP والتسويق محتاجينها.", False),
        ("", False),
        ("إزاي تستخدمه:", True),
        ("• الخانات الصفرا = انت اللي تملاها. الخضرا = اتملت. الرمادي = سيبها (من النظام). الحمرا = محتاجة قرار منك.", False),
        ("• عمود «الحالة» فيه اختيارات: ناقص / اتملى / محتاج مراجعة / مش محتاجينه.", False),
        ("• عمود «الكود (key)» ماتغيّروش — الأجنت بيقرا البيانات بيه.", False),
        ("• لو حاجة مش عارفها سيبها فاضية واكتب في «ملاحظات».", False),
        ("", False),
        ("⚠️ ممنوع تكتب هنا:", True),
        ("• أي باسورد، أو API key، أو Secret، أو رمز تحقق. دول بيتبعتوا لي لوحدهم ويتحطوا على السيرفر بس.", False),
        ("• المسموح: الأسماء، الأرقام، الإيميلات، الروابط، أرقام الحسابات العامة (مثلًا Merchant ID).", False),
        ("", False),
        ("الشيتات:", True),
        ("1. بيانات الشركة — الاسم القانوني، الترخيص، العنوان، التليفونات، ساعات الدعم.", False),
        ("2. المنصات والسوشيال — جوجل، تريب أدفايزر، ميتا، تيك توك، يوتيوب، ياندكس، منصات الحجز…", False),
        ("3. الدومين والسيرفر والخدمات — الدومين، الاستضافة، الإيميل، Paymob، النسخ الاحتياطي.", False),
        ("4. الرحلات — كل رحلة بكل تفاصيلها (آخر نسخة بعد تعديلاتك).", False),
        ("5. خيارات الوحدات — الرحلات اللي سعرها بالوحدة (باجي، مركب، عربية).", False),
        ("6. الصور والميديا — اللوجو وصور كل رحلة وحقوق استخدامها.", False),
        ("7. الترجمة — مين يراجع الروسي والإيطالي.", False),
        ("8. أسئلة مفتوحة — كل سؤال محتاج قرارك، وردّك جنبه.", False),
        ("9. سجل التغييرات — اكتب فيه أي تعديل عملته وتاريخه.", False),
        ("", False),
        ("لما تخلص: ابعت الملف أو قول «حدّثت ملف البيانات» وأنا هقراه وأطبّقه.", True),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = ws.cell(i, 1, text)
        cell.font = Font(name=FONT, bold=bold, size=14 if i == 1 else 11, color="0A2342" if bold else "000000")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def company(wb):
    rows = [
        ("brand_name", "اسم البراند (EN)", "Safari Tours Sharm", "Safari Tours Sharm", "ضروري للإطلاق", "اتملى", ""),
        ("brand_name_ar", "اسم البراند (عربي)", "", "سفاري تورز شرم", "ضروري للإطلاق", "ناقص", ""),
        ("legal_name", "الاسم القانوني للشركة (زي السجل التجاري)", "", "شركة … للسياحة", "ضروري للإطلاق", "ناقص", "بيظهر في الشروط والفواتير وسياسة الخصوصية"),
        ("commercial_register", "رقم السجل التجاري", "", "123456", "ضروري للإطلاق", "ناقص", ""),
        ("tax_id", "الرقم الضريبي", "", "123-456-789", "ضروري للإطلاق", "ناقص", ""),
        ("tourism_license", "رقم ترخيص وزارة السياحة (لو موجود)", "", "", "ضروري للإطلاق", "ناقص", "تريب أدفايزر وجوجل بيثقوا فيه"),
        ("address_ar", "العنوان بالعربي", "", "شارع … خليج نعمة، شرم الشيخ", "ضروري للإطلاق", "ناقص", ""),
        ("address_en", "العنوان بالإنجليزي", "", "… Street, Naama Bay, Sharm El Sheikh", "ضروري للإطلاق", "ناقص", ""),
        ("google_maps_url", "رابط المكتب على خرايط جوجل", "", "https://maps.app.goo.gl/…", "ضروري للإطلاق", "ناقص", ""),
        ("phone_main", "رقم التليفون الأساسي", "+201111292690", "+20…", "ضروري للإطلاق", "اتملى", "من الموقع الحالي — أكّد"),
        ("whatsapp_number", "رقم واتساب الحجز", "+201111292690", "+20…", "ضروري للإطلاق", "اتملى", "أكّد إنه رقم واتساب Business"),
        ("email_public", "الإيميل اللي بيظهر للعملاء", "safaritourssharm@gmail.com", "info@…", "ضروري للإطلاق", "اتملى", "يفضل إيميل على الدومين (info@الدومين)"),
        ("email_bookings", "إيميل بتوصله إشعارات الحجوزات", "", "bookings@…", "ضروري للإطلاق", "ناقص", ""),
        ("support_hours", "ساعات الرد على العملاء", "", "يوميًا 8 ص – 11 م (توقيت مصر)", "ضروري للإطلاق", "ناقص", ""),
        ("support_languages", "لغات الرد على العملاء", "", "عربي، إنجليزي، روسي", "ضروري للإطلاق", "ناقص", ""),
        ("founded_year", "سنة بداية الشغل", "", "2015", "مهم بعد الإطلاق", "ناقص", "بتظهر في «من نحن»"),
        ("owner_name", "اسم صاحب الشركة/المدير (لصفحة من نحن)", "", "", "اختياري", "ناقص", ""),
        ("about_story", "قصة قصيرة عن الشركة (سطرين-تلاتة)", "", "", "مهم بعد الإطلاق", "ناقص", "ممكن تكتبها بالعربي وأنا أترجمها"),
        ("cancellation_policy", "سياسة الإلغاء", "48 ساعة استرداد كامل، 24–48 ساعة 50%، أقل من كده لا يوجد استرداد", "", "ضروري للإطلاق", "اتملى", "دي المعتمدة — غيّرها لو اتغيرت"),
        ("payment_methods", "طرق الدفع المقبولة", "كارت أونلاين (Paymob)", "كارت، كاش عند الاستلام؟", "ضروري للإطلاق", "محتاج مراجعة", "هل في دفع كاش؟"),
    ]
    cols = ["الكود (key)", "البيان", "القيمة", "مثال", "الأولوية", "الحالة", "ملاحظات"]
    sheet(wb, "بيانات الشركة", cols, rows, {"القيمة", "الحالة", "ملاحظات"}, [24, 40, 44, 34, 18, 16, 40], "الحالة", "الأولوية")


def platforms(wb):
    rows = [
        # key, platform, why, exists?, url/handle, login email (no password), what I need, priority, status, notes
        ("google_business", "Google Business Profile (خرايط جوجل)", "أهم مصدر حجوزات وتقييمات في جوجل", "", "", "", "رابط البروفايل + تضيفني Manager بإيميل الشغل", "ضروري للإطلاق", "ناقص", ""),
        ("google_search_console", "Google Search Console", "عشان جوجل يقرا الموقع الجديد ونحوّل الروابط القديمة", "", "", "", "صلاحية Owner/Full على الدومين", "ضروري للإطلاق", "ناقص", ""),
        ("google_analytics", "Google Analytics 4", "قياس الزوار والحجوزات", "", "", "", "Measurement ID (G-…) + صلاحية Editor", "ضروري للإطلاق", "ناقص", "لو مش موجود أعمله أنا"),
        ("google_ads", "Google Ads", "إعلانات جوجل (لو هتعلن)", "", "", "", "رقم الحساب (Customer ID)", "اختياري", "ناقص", ""),
        ("google_review_link", "رابط «اكتب تقييم» في جوجل", "بيتبعت للعميل بعد الرحلة", "", "", "", "الرابط القصير من Google Business", "ضروري للإطلاق", "ناقص", "بيستخدم في إيميل بعد الرحلة"),
        ("tripadvisor", "TripAdvisor", "ثقة السياح الأجانب + تقييمات", "", "", "", "رابط صفحة الشركة + صلاحية Management Center", "ضروري للإطلاق", "ناقص", ""),
        ("tripadvisor_review_link", "رابط تقييم TripAdvisor", "بيتبعت للعميل بعد الرحلة", "", "", "", "رابط Write a review", "مهم بعد الإطلاق", "ناقص", ""),
        ("meta_business", "Meta Business Suite", "إدارة فيسبوك + إنستجرام + واتساب من مكان واحد", "", "", "", "Business ID + تضيفني Admin/Employee", "ضروري للإطلاق", "ناقص", ""),
        ("facebook_page", "Facebook Page", "تواصل وإعلانات", "", "", "", "رابط الصفحة", "ضروري للإطلاق", "ناقص", ""),
        ("instagram", "Instagram", "صور وريلز الرحلات", "", "", "", "اسم الحساب @…", "ضروري للإطلاق", "ناقص", ""),
        ("whatsapp_business", "WhatsApp Business", "الحجز بالطلب والدعم", "آه", "+201111292690", "", "أكّد نوع الحساب (App ولا API) + الكتالوج", "ضروري للإطلاق", "محتاج مراجعة", ""),
        ("meta_pixel", "Meta Pixel / Dataset", "قياس الإعلانات على فيسبوك وإنستجرام", "", "", "", "Pixel ID (رقم)", "مهم بعد الإطلاق", "ناقص", "بيتفعّل بعد موافقة الكوكيز بس"),
        ("tiktok", "TikTok", "فيديوهات قصيرة للرحلات", "", "", "", "اسم الحساب", "اختياري", "ناقص", ""),
        ("youtube", "YouTube", "فيديوهات الرحلات للموقع", "", "", "", "رابط القناة", "اختياري", "ناقص", ""),
        ("yandex_business", "Yandex Business (Яндекс Бизнес)", "السياح الروس بيدوّروا على ياندكس مش جوجل", "", "", "", "رابط الشركة على خرايط ياندكس", "مهم بعد الإطلاق", "ناقص", "مهم جدًا للسوق الروسي"),
        ("yandex_webmaster", "Yandex Webmaster + Metrica", "ظهور الموقع في ياندكس", "", "", "", "صلاحية على الدومين", "مهم بعد الإطلاق", "ناقص", ""),
        ("vk", "VK (ВКонтакте)", "سوشيال السوق الروسي", "", "", "", "رابط الصفحة", "اختياري", "ناقص", ""),
        ("telegram", "Telegram", "تواصل مع السياح الروس", "", "", "", "اسم القناة أو الحساب", "اختياري", "ناقص", ""),
        ("bing_webmaster", "Bing Webmaster", "ظهور في Bing", "", "", "", "صلاحية على الدومين", "اختياري", "ناقص", "ممكن أعمله أنا من Search Console"),
        ("apple_business", "Apple Business Connect", "الظهور في خرايط آيفون", "", "", "", "رابط أو ID", "اختياري", "ناقص", ""),
        ("viator", "Viator", "منصة بيع رحلات (عمولة)", "", "", "", "رابط صفحة المورد (Supplier)", "اختياري", "ناقص", "لو بتبيع عليها"),
        ("getyourguide", "GetYourGuide", "منصة بيع رحلات (عمولة)", "", "", "", "رابط صفحة المورد", "اختياري", "ناقص", "لو بتبيع عليها"),
        ("trustpilot", "Trustpilot", "تقييمات للأوروبيين", "", "", "", "رابط الصفحة", "اختياري", "ناقص", ""),
    ]
    cols = ["الكود (key)", "المنصة", "ليه محتاجينها", "عندك حساب؟ (آه/لأ)", "الرابط أو اسم الحساب", "إيميل الحساب (من غير باسورد)", "المطلوب منك بالظبط", "الأولوية", "الحالة", "ملاحظات"]
    ws = sheet(wb, "المنصات والسوشيال", cols, rows, {"عندك حساب؟ (آه/لأ)", "الرابط أو اسم الحساب", "إيميل الحساب (من غير باسورد)", "الحالة", "ملاحظات"}, [22, 30, 34, 14, 36, 30, 40, 18, 16, 30], "الحالة", "الأولوية")
    dv = DataValidation(type="list", formula1='"آه,لأ,مش عارف"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D2:D80")


def services(wb):
    rows = [
        ("domain_name", "الدومين الحالي", "safaritourssharm.com", "", "ضروري للإطلاق", "اتملى", "أكّد"),
        ("domain_registrar", "الدومين متسجل فين؟", "", "GoDaddy / Namecheap / Hostinger…", "ضروري للإطلاق", "ناقص", ""),
        ("dns_access", "مين معاه صلاحية الـDNS؟", "", "انت / مبرمج قديم…", "ضروري للإطلاق", "ناقص", "هنحتاجه يوم الإطلاق بس"),
        ("current_hosting", "الموقع القديم (WordPress) مستضاف فين؟", "", "", "ضروري للإطلاق", "ناقص", "عشان نحوّل الروابط القديمة"),
        ("server_choice", "السيرفر الجديد", "", "VPS عند Hostinger / سيرفر موجود", "ضروري للإطلاق", "ناقص", "أقدر أرشّحلك لو مش محدد"),
        ("smtp_provider", "مزوّد الإيميلات (رسايل التأكيد)", "", "Brevo / Postmark / Amazon SES / Zoho", "ضروري للإطلاق", "ناقص", "أرشّح Brevo أو Postmark"),
        ("sending_domain", "الدومين اللي الإيميلات هتطلع منه", "", "bookings@safaritourssharm.com", "ضروري للإطلاق", "ناقص", ""),
        ("paymob_account", "حساب Paymob موجود؟", "", "آه / لأ", "ضروري للإطلاق", "ناقص", ""),
        ("paymob_merchant_id", "Paymob Merchant ID (رقم عام)", "", "123456", "ضروري للإطلاق", "ناقص", "المفاتيح السرية تتبعت لوحدها — مش هنا"),
        ("paymob_eur_enabled", "Paymob مفعّل الدفع باليورو؟", "", "آه / لأ", "ضروري للإطلاق", "ناقص", "لازم يتأكد مع Paymob"),
        ("backup_location", "النسخ الاحتياطي يتحفظ فين؟", "", "Google Drive / Backblaze…", "مهم بعد الإطلاق", "ناقص", ""),
        ("cloudflare", "تستخدم Cloudflare؟", "", "آه / لأ", "اختياري", "ناقص", ""),
        ("media_drive_link", "رابط فولدر الصور والفيديوهات", "", "Google Drive link", "ضروري للإطلاق", "ناقص", "لصور الرحلات واللوجو"),
    ]
    cols = ["الكود (key)", "البيان", "القيمة", "مثال", "الأولوية", "الحالة", "ملاحظات"]
    sheet(wb, "الدومين والسيرفر والخدمات", cols, rows, {"القيمة", "الحالة", "ملاحظات"}, [24, 40, 40, 36, 18, 16, 40], "الحالة", "الأولوية")


SLOT_AR = {"SUNRISE": "شروق", "MORNING": "صباحًا", "AFTERNOON": "بعد الظهر", "SUNSET": "غروب"}
TYPE_AR = {"TOUR": "رحلة", "REQUEST_ONLY": "بالطلب", "TRANSFER": "انتقال"}
UNIT_TOURS = {"double-buggy-camel-ride", "speed-boat-adventure", "sharm-airport-transfer"}


def tours(wb):
    """Owner's latest trips sheet, with the approved revisions applied on top."""
    source = openpyxl.load_workbook(RESEARCH / "drafts" / "owner-edited-2026-09-30.xlsx")["الرحلات"]
    header = [c.value for c in source[1]]
    rows = [list(r) for r in source.iter_rows(min_row=2, values_only=True) if r and r[0]]
    col = {name: i for i, name in enumerate(header)}
    catalog = json.load(open(RESEARCH / "approved-catalog.json"))
    to_column = {
        "priceAdultEur": "سعر البالغ €",
        "durationText": "المدة (EN)",
        "pricingNote": "ملاحظة السعر (EN)",
        "capacity": "السعة (أفراد في الميعاد)",
    }
    by_slug = {r[0]: r for r in rows}
    for revision in catalog.get("revisions", []):
        for change in revision["changes"]:
            row = by_slug.get(change["slug"])
            if row is None:
                continue
            field, value = change["field"], change["to"]
            if field in to_column:
                row[col[to_column[field]]] = value
            elif field == "availableTimeSlots":
                row[col["الأوقات المتاحة"]] = "، ".join(SLOT_AR[s] for s in value)
            elif field == "tourType":
                row[col["النوع"]] = TYPE_AR[value]
            elif field == "isActive":
                row[col["ظاهرة للبيع؟"]] = "آه" if value else "لأ"
    # Seats confirmed by the owner on 2026-10-01.
    for slug, capacity in (("speed-boat-adventure", 5), ("sharm-airport-transfer", 8)):
        if slug in by_slug:
            by_slug[slug][col["السعة (أفراد في الميعاد)"]] = capacity
    out_header = header[:]
    insert_at = header.index("سعر البالغ €")
    out_header.insert(insert_at, "طريقة التسعير")
    out_rows = []
    for row in rows:
        basis = "بالوحدة (شوف شيت خيارات الوحدات)" if row[0] in UNIT_TOURS else "بالفرد"
        out_rows.append(row[:insert_at] + [basis] + row[insert_at:])
    locked = {out_header[0], "طريقة التسعير", "⚠️ أسئلة محتاجة قرارك"}
    editable = set(out_header) - locked
    widths = [26 if i == 0 else 18 for i in range(len(out_header))]
    sheet(wb, "الرحلات", out_header, out_rows, editable, widths)


def unit_options(wb):
    rows = [
        ("double-buggy-camel-ride", "buggy", "باجي بمقعدين", "Two-seat buggy", 2, 30, "اتملى", "سعة الميعاد 10 = 5 باجي"),
        ("speed-boat-adventure", "boat", "مركب سريع خاص", "Private speedboat", 5, 150, "اتملى", "مركب واحد في الميعاد"),
        ("sharm-airport-transfer", "sedan", "سيارة سيدان", "Sedan car", 4, 15, "اتملى", ""),
        ("sharm-airport-transfer", "suv", "سيارة SUV", "SUV", 4, 20, "اتملى", ""),
        ("sharm-airport-transfer", "minibus", "ميني باص هاي إس", "Hiace minibus", 8, 35, "اتملى", "سعة الميعاد 8"),
    ]
    cols = ["كود الرحلة", "كود الوحدة", "الاسم بالعربي", "الاسم بالإنجليزي", "عدد الأفراد في الوحدة", "السعر € للوحدة", "الحالة", "ملاحظات"]
    sheet(wb, "خيارات الوحدات", cols, rows, {"الاسم بالعربي", "الاسم بالإنجليزي", "عدد الأفراد في الوحدة", "السعر € للوحدة", "الحالة", "ملاحظات"}, [26, 14, 22, 22, 16, 14, 14, 30], "الحالة")


def media(wb):
    catalog = json.load(open(RESEARCH / "approved-catalog.json"))
    rows = [
        ("logo", "اللوجو الأساسي (ملوّن)", "SVG أو PNG بخلفية شفافة", "", "", "", "ناقص", ""),
        ("logo_white", "اللوجو الأبيض (للخلفيات الغامقة)", "SVG أو PNG", "", "", "", "ناقص", ""),
        ("brand_colors", "ألوان البراند (لو عندك)", "أكواد زي #F28C28", "", "", "", "ناقص", "لو مفيش أنا مختارها"),
        ("hero_video", "فيديو قصير للصفحة الرئيسية (5–8 ثواني)", "MP4 أفقي، من غير صوت", "", "", "", "ناقص", "اختياري بس بيفرق جدًا"),
    ]
    for tour in catalog["tours"]:
        rows.append((tour["slug"], f"صور: {tour.get('nameEn') or tour['slug']}", "صورة غلاف + 4–8 صور (أفقي، 1600px أو أكبر)", "", "", "", "ناقص", ""))
    cols = ["الكود (key)", "المطلوب", "المواصفات", "رابط الملفات (Drive)", "حقوق الاستخدام مضمونة؟ (آه/لأ)", "مين صوّرها", "الحالة", "ملاحظات"]
    sheet(wb, "الصور والميديا", cols, rows, {"رابط الملفات (Drive)", "حقوق الاستخدام مضمونة؟ (آه/لأ)", "مين صوّرها", "الحالة", "ملاحظات"}, [28, 40, 36, 36, 18, 18, 14, 30], "الحالة")


def translation(wb):
    rows = [
        ("en", "الإنجليزي", "اتكتب واتعتمد منك", "", "اتملى", ""),
        ("ar", "العربي", "أنا هترجمه من الإنجليزي", "", "ناقص", "هتراجعه بنفسك"),
        ("ru", "الروسي", "أنا هترجمه — محتاج متحدث روسي يراجعه", "", "ناقص", "اسم ورقم المراجع"),
        ("it", "الإيطالي", "أنا هترجمه — محتاج متحدث إيطالي يراجعه", "", "ناقص", "اسم ورقم المراجع"),
    ]
    cols = ["الكود (key)", "اللغة", "الخطة", "اسم المراجع", "الحالة", "ملاحظات"]
    sheet(wb, "الترجمة", cols, rows, {"اسم المراجع", "الحالة", "ملاحظات"}, [12, 14, 44, 28, 14, 36], "الحالة")


def questions(wb):
    source = openpyxl.load_workbook(RESEARCH / "drafts" / "owner-edited-2026-09-30.xlsx")["الرحلات"]
    header = [c.value for c in source[1]]
    q = header.index("⚠️ أسئلة محتاجة قرارك")
    a = header.index("✍️ ردك / ملاحظاتك")
    answered = {
        "ras-mohamed-white-island-boat": ("إيجار معدات السنوركل بفلوس إضافية (رد 2026-10-01).", "اتملى"),
        "bedouin-dinner-camel-ride": ("ردك 2026-10-01: «العشاء البدوي في الخاص في اي بي انكلوديت». فهمي: الخيمة الخاصة VIP مشمولة في العشاء البدوي — أكّد الصياغة.", "محتاج مراجعة"),
        "sharm-airport-transfer": ("سيدان 4 أفراد €15 · SUV 4 أفراد €20 · هاي إس 8 أفراد €35 — اتأكد 2026-10-01.", "اتملى"),
        "speed-boat-adventure": ("مركب خاص €150 لحد 5 أفراد، المدة 3 ساعات — اتأكد 2026-10-01.", "اتملى"),
    }
    rows = []
    for r in source.iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or not r[q]:
            continue
        reply = r[a] or ""
        status = "محتاج مراجعة" if reply else "ناقص"
        if r[0] in answered:
            reply, status = answered[r[0]]
        rows.append((f"q-{r[0]}", r[0], str(r[q]), reply, status))
    rows.append(("q-snorkel-equipment", "عام", "إيجار معدات السنوركل في رحلات البحر بفلوس ولا مشمول؟", "بفلوس إضافية (رد 2026-10-01).", "اتملى"))
    cols = ["الكود (key)", "الرحلة", "السؤال", "ردّك", "الحالة"]
    sheet(wb, "أسئلة مفتوحة", cols, rows, {"ردّك", "الحالة"}, [30, 26, 70, 50, 16], "الحالة")


def changelog(wb):
    rows = [
        (date.today().isoformat(), "Claude", "إنشاء الملف وتجميع كل البيانات المطلوبة في شيتات"),
        ("2026-10-01", "محمد", "تأكيد سعة SUV (4) والـSpeed Boat (5) وميعاد التوصيل (8)؛ معدات السنوركل بفلوس"),
    ]
    sheet(wb, "سجل التغييرات", ["التاريخ", "مين", "إيه اللي اتغيّر"], rows, {"التاريخ", "مين", "إيه اللي اتغيّر"}, [16, 16, 90])


def parse_args():
    parser = argparse.ArgumentParser(
        description="Create the original Safari Tours owner-data bootstrap workbook."
    )
    parser.add_argument(
        "--force-bootstrap",
        action="store_true",
        help="overwrite the live workbook with the original bootstrap (data loss risk)",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    if OUT.exists() and not args.force_bootstrap:
        raise SystemExit(
            f"Refusing to overwrite the live owner workbook: {OUT}\n"
            "The workbook contains owner answers and audited account data. "
            "Use --force-bootstrap only when intentionally creating a disposable base copy."
        )
    wb = Workbook()
    readme(wb)
    company(wb)
    platforms(wb)
    services(wb)
    tours(wb)
    unit_options(wb)
    media(wb)
    translation(wb)
    questions(wb)
    changelog(wb)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
