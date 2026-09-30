"""Builds the owner's editable workbook of every Safari tour: approved catalog
facts + EN content drafts + facts drafts + review flags. Read back with
read_workbook.py after the owner edits it."""
import json
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

catalog = json.load(open("../approved-catalog.json"))
drafts = {t["slug"]: t for t in json.load(open("tour-content-drafts.en.json"))["tours"]}

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="0A2342")
EDIT = PatternFill("solid", fgColor="FFF8DC")     # editable: light yellow
LOCK = PatternFill("solid", fgColor="E7E6E6")     # do not edit: grey
FLAG = PatternFill("solid", fgColor="FDE9E7")     # review questions: light red
DECIDE = PatternFill("solid", fgColor="E2F0D9")   # owner answers: light green
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CATEGORIES = {"DESERT": "صحرا", "SEA": "بحر", "CULTURAL": "ثقافة", "SHOWS": "عروض", "TRANSFERS": "انتقالات"}
SLOTS = {"SUNRISE": "شروق", "MORNING": "صباحًا", "AFTERNOON": "بعد الظهر", "SUNSET": "غروب"}
PICKUP = {"INCLUDED": "مشمول", "NOT_INCLUDED": "مش مشمول", "SOME_AREAS": "بعض المناطق", None: ""}

def yn(value):
    return {True: "آه", False: "لأ", None: ""}[value]

# (key, header, width, fill, note)
COLUMNS = [
    ("slug", "الكود (slug)\nماتغيّروش", 26, LOCK, "المعرّف الثابت للرحلة في النظام والرابط. ماينفعش يتغيّر."),
    ("active", "ظاهرة للبيع؟", 11, EDIT, "آه = تظهر على الموقع وتتحجز. لأ = مخفية."),
    ("name", "اسم الرحلة (EN)", 28, EDIT, None),
    ("category", "الفئة", 11, EDIT, "صحرا / بحر / ثقافة / عروض / انتقالات"),
    ("type", "النوع", 12, EDIT, "رحلة = تتحجز وتتدفع. بالطلب = سعر على الطلب (واتساب). انتقال = مواصلات."),
    ("duration", "المدة (EN)", 20, EDIT, "زي ما هتظهر للعميل بالإنجليزي، مثال: 5–6 hours"),
    ("price_adult", "سعر البالغ €", 11, EDIT, "رقم بس باليورو. فاضي = سعر على الطلب."),
    ("price_child", "سعر الطفل €", 11, EDIT, "فاضي = مفيش سعر أطفال مختلف."),
    ("price_note", "ملاحظة السعر (EN)", 26, EDIT, "مثال: Price per buggy (seats 2 adults)"),
    ("slots", "الأوقات المتاحة", 20, EDIT, "افصل بفاصلة: شروق، صباحًا، بعد الظهر، غروب"),
    ("days", "أيام التشغيل", 18, EDIT, "لو الرحلة مش يومية: مثال «الأحد، الخميس». فاضي = كل يوم."),
    ("capacity", "السعة (أفراد في الميعاد)", 12, EDIT, None),
    ("short", "وصف مختصر (EN)\nسطر أو سطرين", 45, EDIT, "بيظهر على كارت الرحلة. حد أقصى 300 حرف."),
    ("description", "الوصف الكامل (EN)", 70, EDIT, "حد أقصى 6000 حرف."),
    ("includes", "المشمول (EN)\nكل بند في سطر", 42, EDIT, "كل بند في سطر لوحده (Alt+Enter في إكسيل)."),
    ("excludes", "غير المشمول (EN)\nكل بند في سطر", 34, EDIT, None),
    ("know", "اعرف قبل ما تروح (EN)\nكل بند في سطر", 48, EDIT, "هات إيه، مش مناسب لمين، ملاحظات مهمة."),
    ("meeting", "نقطة التجمع / الاستلام (EN)", 34, EDIT, None),
    ("children", "مناسب للأطفال؟", 11, EDIT, "آه / لأ / فاضي لو مش متأكد"),
    ("min_age", "أقل سن", 9, EDIT, "رقم، أو فاضي"),
    ("pickup", "الاستلام من الفندق", 14, EDIT, "مشمول / مش مشمول / بعض المناطق"),
    ("guide_langs", "لغات المرشد", 14, EDIT, "أكواد: en, ru, it, ar … مفصولة بفاصلة"),
    ("flags", "⚠️ أسئلة محتاجة قرارك", 60, FLAG, "النقط اللي لقيتها متعارضة أو مش مؤكدة في الموقع القديم."),
    ("answer", "✍️ ردك / ملاحظاتك", 45, DECIDE, "اكتب هنا ردك على الأسئلة أو أي تعديل عايزه."),
    ("status", "الحالة", 12, DECIDE, "مسودة = لسه. معتمد = راجعتها وموافق تتنشر."),
]

wb = Workbook()

# ── Instructions ────────────────────────────────────────────────────────────
guide = wb.active
guide.title = "تعليمات"
guide.sheet_view.rightToLeft = True
lines = [
    ("ملف رحلات Safari Tours Sharm — كل البيانات في مكان واحد", True),
    ("", False),
    ("إزاي تستخدم الملف:", True),
    ("• كل صف = رحلة واحدة (30 رحلة) في شيت «الرحلات».", False),
    ("• الخانات الصفرا = تعدّل فيها براحتك.", False),
    ("• الخانات الرمادي (الكود) = ماتغيّرهاش، ده رابط الرحلة في النظام.", False),
    ("• العمود الأحمر الفاتح = أسئلة لقيتها متعارضة أو مش مؤكدة في الموقع القديم.", False),
    ("• العمود الأخضر = اكتب ردك فيه، وفي الآخر غيّر «الحالة» لـ«معتمد» لما توافق على الرحلة.", False),
    ("• القوايم (المشمول، غير المشمول، اعرف قبل ما تروح): كل بند في سطر لوحده — في إكسيل اضغط Alt+Enter لسطر جديد.", False),
    ("• النصوص للعملاء بالإنجليزي. العربي والروسي والإيطالي هنترجمهم بعد ما الإنجليزي يتعتمد.", False),
    ("• لو في حاجة مش متأكد منها، سيب الخانة فاضية — الموقع مش هيعرض حاجة مش مؤكدة.", False),
    ("", False),
    ("سياسة الإلغاء المعتمدة حاليًا (لكل الرحلات):", True),
    ("• استرداد كامل قبل الرحلة بـ48 ساعة أو أكتر.", False),
    ("• 50% من 24 لـ48 ساعة.", False),
    ("• أقل من 24 ساعة أو عدم الحضور: مفيش استرداد.", False),
    ("لو عايز سياسة مختلفة لرحلة معينة، اكتبها في عمود «ردك».", False),
    ("", False),
    ("لما تخلص: ابعت الملف تاني، وأنا أقراه وأحدّث النظام، وأوريك الفرق قبل أي نشر.", True),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = guide.cell(row=i, column=1, value=text)
    cell.font = Font(name=FONT, size=14 if i == 1 else 11, bold=bold, color="0A2342" if bold else "000000")
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
guide.column_dimensions["A"].width = 110
legend = [("خانة تعدّل فيها", EDIT), ("ماتعدّلش", LOCK), ("أسئلة محتاجة قرارك", FLAG), ("ردك والحالة", DECIDE)]
start = len(lines) + 2
guide.cell(row=start, column=1, value="دليل الألوان:").font = Font(name=FONT, bold=True, color="0A2342")
for k, (label, fill) in enumerate(legend, start=1):
    c = guide.cell(row=start + k, column=1, value=label)
    c.fill = fill
    c.font = Font(name=FONT)
    c.border = BORDER
    c.alignment = Alignment(horizontal="right", readingOrder=2)

# ── Tours ───────────────────────────────────────────────────────────────────
ws = wb.create_sheet("الرحلات")
ws.sheet_view.rightToLeft = True
for col, (key, header, width, fill, note) in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = HEAD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    cell.border = BORDER
    if note:
        cell.comment = Comment(note, "Claude")
    ws.column_dimensions[cell.column_letter].width = width
ws.row_dimensions[1].height = 46

TYPE = {"TOUR": "رحلة", "REQUEST_ONLY": "بالطلب", "TRANSFER": "انتقال"}
for r, tour in enumerate(catalog["tours"], start=2):
    d = drafts[tour["slug"]]
    c, f = d["content"], d.get("facts", {})
    values = {
        "slug": tour["slug"],
        "active": yn(tour["isActive"]),
        "name": c["name"],
        "category": CATEGORIES[tour["category"]],
        "type": TYPE[tour["tourType"]],
        "duration": tour["durationText"],
        "price_adult": tour["priceAdultEur"],
        "price_child": tour.get("priceChildEur"),
        "price_note": tour.get("pricingNote") or "",
        "slots": "، ".join(SLOTS[s] for s in tour["availableTimeSlots"]),
        "days": "",
        "capacity": tour["capacity"],
        "short": c["shortDescription"],
        "description": c["description"],
        "includes": "\n".join(c["includes"]),
        "excludes": "\n".join(c["excludes"]),
        "know": "\n".join(c["knowBeforeYouGo"]),
        "meeting": c.get("meetingPoint") or "",
        "children": yn(f.get("childrenAllowed")),
        "min_age": f.get("minimumAge"),
        "pickup": PICKUP[f.get("hotelPickup")],
        "guide_langs": ", ".join(f.get("guideLanguages", [])),
        "flags": "\n".join(f"{i}) {q}" for i, q in enumerate(d["reviewFlags"], start=1)),
        "answer": "",
        "status": "مسودة",
    }
    for col, (key, _h, _w, fill, _n) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=r, column=col, value=values[key])
        cell.fill = fill
        cell.border = BORDER
        cell.font = Font(name=FONT, size=10, bold=key in ("slug", "name"))
        rtl = key in ("flags", "answer", "status", "active", "category", "type", "slots", "days", "children", "pickup")
        cell.alignment = Alignment(
            horizontal="right" if rtl else "left",
            vertical="top",
            wrap_text=True,
            readingOrder=2 if rtl else 1,
        )
        if key in ("price_adult", "price_child"):
            cell.number_format = '#,##0.00 "€"'
    ws.row_dimensions[r].height = 190

last = len(catalog["tours"]) + 1
col_of = {key: ws.cell(row=1, column=i).column_letter for i, (key, *_rest) in enumerate(COLUMNS, start=1)}
def dropdown(key, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = "اختار من القايمة"
    ws.add_data_validation(dv)
    dv.add(f"{col_of[key]}2:{col_of[key]}{last}")
dropdown("active", ["آه", "لأ"])
dropdown("category", list(CATEGORIES.values()))
dropdown("type", list(TYPE.values()))
dropdown("children", ["آه", "لأ"])
dropdown("pickup", ["مشمول", "مش مشمول", "بعض المناطق"])
dropdown("status", ["مسودة", "معتمد"])
for key in ("price_adult", "price_child", "capacity", "min_age"):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.error = "رقم بس"
    ws.add_data_validation(dv)
    dv.add(f"{col_of[key]}2:{col_of[key]}{last}")

ws.freeze_panes = "D2"   # keep code, active flag and name visible while scrolling
ws.auto_filter.ref = f"A1:{col_of['status']}{last}"

# ── Summary (live counts) ───────────────────────────────────────────────────
summary = wb.create_sheet("ملخص")
summary.sheet_view.rightToLeft = True
rows = [
    ("عدد الرحلات", f"=COUNTA('الرحلات'!A2:A{last})"),
    ("معتمدة", f"=COUNTIF('الرحلات'!{col_of['status']}2:{col_of['status']}{last},\"معتمد\")"),
    ("لسه مسودة", f"=COUNTIF('الرحلات'!{col_of['status']}2:{col_of['status']}{last},\"مسودة\")"),
    ("رحلات فيها أسئلة", f"=COUNTIF('الرحلات'!{col_of['flags']}2:{col_of['flags']}{last},\"?*\")"),
    ("رحلات رديت عليها", f"=COUNTIF('الرحلات'!{col_of['answer']}2:{col_of['answer']}{last},\"?*\")"),
    ("ظاهرة للبيع", f"=COUNTIF('الرحلات'!{col_of['active']}2:{col_of['active']}{last},\"آه\")"),
]
for i, (label, formula) in enumerate(rows, start=1):
    a = summary.cell(row=i, column=1, value=label)
    b = summary.cell(row=i, column=2, value=formula)
    for cell in (a, b):
        cell.font = Font(name=FONT, size=12, bold=cell is a)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="right", readingOrder=2)
summary.column_dimensions["A"].width = 26
summary.column_dimensions["B"].width = 12

wb.active = wb.index(ws)
wb.save("safari-tours-all-trips.xlsx")
print("saved", last - 1, "tours")
